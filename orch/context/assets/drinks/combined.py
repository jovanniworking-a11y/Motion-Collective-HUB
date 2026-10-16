"""Devil's Night · combined budget + shopping workbook (5 Oct 2026)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

A = "Arial"
def f(**k): return Font(name=A, size=k.pop("size", 10), **k)
BLUE, GREEN, GREY = "0000FF", "008000", "6B7280"
YEL = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="1F1F24")
SUB = PatternFill("solid", fgColor="EEF1F6")
thin = Side(style="thin", color="D0D4DC")
USD, INT, PCT, D1 = '$#,##0.00;($#,##0.00);"-"', '#,##0;(#,##0);"-"', '0.0%', '0.0'
TW, HEB, WM, SAMS = "Twin Liquors", "H-E-B", "Walmart", "Sam's Club"

wb = Workbook()

def sheet(name, widths):
    s = wb.create_sheet(name)
    for i, w in enumerate(widths, 1): s.column_dimensions[L(i)].width = w
    s.sheet_view.showGridLines = False
    return s

def title(s, t, sub):
    s["A1"] = t; s["A1"].font = f(size=14, bold=True)
    s["A2"] = sub; s["A2"].font = f(italic=True, color=GREY)

def header(s, r, cols):
    for i, c in enumerate(cols, 1):
        x = s.cell(r, i, c); x.font = f(bold=True, color="FFFFFF"); x.fill = HEAD
        x.alignment = Alignment(wrap_text=True, vertical="center")
    s.row_dimensions[r].height = 30

def put(s, r, c, v, fmt=None, inp=False, link=False, bold=False, wrap=False, fill=None):
    x = s.cell(r, c, v)
    x.font = f(color=BLUE if inp else GREEN if link else "000000", bold=bold)
    if fmt: x.number_format = fmt
    if wrap: x.alignment = Alignment(wrap_text=True, vertical="top")
    if fill: x.fill = fill
    x.border = Border(bottom=thin)
    return x

def yn(s, rng):
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    s.add_data_validation(dv); dv.add(rng)

# ---------------------------------------------------------------- Drinks
d = sheet("Drinks", [30, 14, 16, 7, 11, 12, 9, 8, 10, 8, 46])
title(d, "Drinks · 2 jungle juices + cans", "Blue = inputs you can change. Std drink = 0.6 oz alcohol. BYOB tops up the rest.")
header(d, 4, ["Item", "Group", "Store", "Qty", "Unit $", "Cost", "Oz each", "ABV", "Std drinks", "Guest?", "Note / link"])
D = [
 ("Svedka vodka 1.75 L", "Vodka JJ", TW, 1, 18.99, 59.2, .40, "Y", "https://twinliquors.com/shop/product/svedka-80-proof-vodka/564bd7e469702d5e6e6e0000"),
 ("Hawaiian Punch 1 gal", "Vodka JJ", HEB, 1, 3.22, 128, 0, "Y", "H-E-B"),
 ("Pineapple juice 64 oz", "Vodka JJ", HEB, 1, 4.27, 64, 0, "Y", "H-E-B"),
 ("Sprite 2 L", "Vodka JJ", HEB, 1, 2.97, 67.6, 0, "Y", "H-E-B"),
 ("Lunazul Blanco tequila 1.75 L", "Tequila JJ", TW, 1, 36.99, 59.2, .40, "Y", "https://twinliquors.com/ (Lunazul 1.75 L)"),
 ("DeKuyper Triple Sec 1 L (16 oz used)", "Tequila JJ", TW, 1, 7.99, 16, .15, "Y", "Only 16 oz goes in; rest kept."),
 ("H-E-B limeade 52 oz", "Tequila JJ", HEB, 2, 2.48, 52, 0, "Y", "H-E-B"),
 ("H-E-B cranberry cocktail 96 oz", "Tequila JJ", HEB, 1, 4.28, 96, 0, "Y", "https://www.heb.com/product-detail/h-e-b-cranberry-cocktail-juice/5785031"),
 ("Sprite 2 L", "Tequila JJ", HEB, 1, 2.97, 67.6, 0, "Y", "H-E-B"),
 ("BeatBox Party Box 6 × 500 ml", "Cans", HEB, 6, 20.98/6, 16.9, .111, "Y", "https://www.heb.com/product-detail/beatbox-party-box-variety-pack/5508020"),
 ("BuzzBallz Chiller 187 ml", "Cans", HEB, 12, 3.12, 6.3, .15, "Y", "https://www.heb.com/product-detail/buzzballz-sour-apple-chiller/2787503 (ABV assumed 15%)"),
 ("White Claw variety 24 pk", "Cans", HEB, 24, 29.87/24, 12, .05, "Y", "https://www.heb.com/product-detail/white-claw-hard-seltzer-variety-24-pk-cans/3758568"),
 ("Miller Lite 30 pk", "Cans", WM, 30, 24.67/30, 12, .042, "Y", "https://www.walmart.com/ip/10984520"),
 ("Jinro Chamisul Fresh soju 750 ml", "he/they rider", TW, 1, 11.99, 25.4, .165, "N", "For the b2b, not guests."),
]
r0 = 5
for i, (n, g, st, q, u, oz, abv, gu, note) in enumerate(D):
    r = r0 + i
    put(d, r, 1, n); put(d, r, 2, g); put(d, r, 3, st)
    put(d, r, 4, q, INT, inp=True); put(d, r, 5, round(u, 4), USD, inp=True)
    put(d, r, 6, f"=D{r}*E{r}", USD); put(d, r, 7, oz, D1, inp=True); put(d, r, 8, abv, '0.0%;;"-"', inp=True)
    put(d, r, 9, f"=D{r}*G{r}*H{r}/0.6", D1); put(d, r, 10, gu, inp=True); put(d, r, 11, note, wrap=False)
rl = r0 + len(D) - 1
yn(d, f"J{r0}:J{rl}")
# triple sec: only 16 oz used — Qty 1 bottle, oz column 16 (std counts what goes in)
tr = rl + 1
put(d, tr, 1, "Drinks total", bold=True, fill=SUB); put(d, tr, 6, f"=SUM(F{r0}:F{rl})", USD, bold=True, fill=SUB)
put(d, tr, 9, f'=SUMIF(J{r0}:J{rl},"Y",I{r0}:I{rl})', D1, bold=True, fill=SUB)
for c in (2,3,4,5,7,8,10,11): d.cell(tr, c).fill = SUB
DR_TOTAL = f"Drinks!F{tr}"
# batch sizes
b = tr + 2
put(d, b, 1, "Batch check", bold=True)
put(d, b+1, 1, "Vodka JJ volume (gal)"); put(d, b+1, 2, f'=SUMPRODUCT((B{r0}:B{rl}="Vodka JJ")*D{r0}:D{rl}*G{r0}:G{rl})/128', '0.00')
put(d, b+2, 1, "Vodka JJ strength"); put(d, b+2, 2, f'=SUMPRODUCT((B{r0}:B{rl}="Vodka JJ")*D{r0}:D{rl}*G{r0}:G{rl}*H{r0}:H{rl})/(B{b+1}*128)', PCT)
put(d, b+3, 1, "Tequila JJ volume (gal)"); put(d, b+3, 2, f'=SUMPRODUCT((B{r0}:B{rl}="Tequila JJ")*D{r0}:D{rl}*G{r0}:G{rl})/128', '0.00')
put(d, b+4, 1, "Tequila JJ strength"); put(d, b+4, 2, f'=SUMPRODUCT((B{r0}:B{rl}="Tequila JJ")*D{r0}:D{rl}*G{r0}:G{rl}*H{r0}:H{rl})/(B{b+3}*128)', PCT)
put(d, b+5, 1, "Serve: 1 std drink of JJ ≈ oz"); put(d, b+5, 2, f"=0.6/AVERAGE(B{b+2},B{b+4})", D1)
d.cell(b+5, 3, "→ an 8 oz pour in a 16 oz cup over ice ≈ 1 drink").font = f(italic=True, color=GREY)
# need
n = b + 7
put(d, n, 1, "Do we have enough?", bold=True)
put(d, n+1, 1, "Drinking guests"); put(d, n+1, 2, 55, INT, inp=True, fill=YEL)
d.cell(n+1, 3, "70 cap; BYOB, so not everyone drinks ours").font = f(italic=True, color=GREY)
put(d, n+2, 1, "Drinks each from us"); put(d, n+2, 2, 3, INT, inp=True, fill=YEL)
put(d, n+3, 1, "Needed"); put(d, n+3, 2, f"=B{n+1}*B{n+2}", INT)
put(d, n+4, 1, "We have"); put(d, n+4, 2, f"=I{tr}", D1)
put(d, n+5, 1, "Covered", bold=True); put(d, n+5, 2, f"=I{tr}/B{n+3}", PCT, bold=True)
put(d, n+6, 1, "Cost per std drink"); put(d, n+6, 2, f'=SUMIF(J{r0}:J{rl},"Y",F{r0}:F{rl})/I{tr}', USD)
d.cell(n+8, 1, "Backup if JJ runs dry: second vodka batch ≈ $29, second tequila batch ≈ $49 (same recipe).").font = f(italic=True, color=GREY)
d.cell(n+9, 1, "Texas: spirits only at package stores (Twin, Spec's, WB Liquors at Costco Kelly Ln), Mon–Sat 10am–9pm, closed Sun. Buy liquor by Thu 29 Oct; grocery stuff any day.").font = f(italic=True, color=GREY)
DR_NEED = n

# ---------------------------------------------------------------- Ice & Cups
ic = sheet("Ice & Cups", [40, 14, 12, 11, 11, 12, 12, 44])
title(ic, "Ice, cups, serving", "Prices seen 5 Oct 2026. Pick = the one that goes into the budget (Y/N).")
put(ic, 4, 1, "Ice needed (lb)", bold=True); put(ic, 4, 2, 160, INT, inp=True, fill=YEL)
ic.cell(4, 3, "≈ 1 lb per guest for cups + ~90 lb for the cooler/tubs over 5 hours").font = f(italic=True, color=GREY)
header(ic, 6, ["Ice source", "Where", "lb / bag", "$ / bag", "$ / lb", "Bags", "Cost", "Pick?  ·  note"])
ICE = [
 ("Texas Cold Front ice house (Twice the Ice)", "448 FM 685 / 15803 Windermere Dr, Pflugerville · 24/7", 20, 3.00, "Y", "EST $3/20 lb (typical TTI price) — check with free first bag, code FREEICE. tticlub.com/TexasColdFront"),
 ("Walmart bagged ice", "Pflugerville Walmart", 10, 1.88, "N", "Confirmed price. Backup."),
 ("Twin Liquors 20 lb", "1700 E Palm Valley, Round Rock", 20, 4.49, "N", "twinliquors.com"),
 ("H-E-B bagged ice", "any H-E-B", 10, 2.50, "N", ""),
 ("Artic Refresh kiosk", "several Austin kiosks · 24/7", 10, 3.00, "N", "articrefresh.com"),
]
for i, (nm, wh, lb, p, pk, note) in enumerate(ICE):
    r = 7 + i
    put(ic, r, 1, nm); put(ic, r, 2, wh); put(ic, r, 3, lb, INT, inp=True); put(ic, r, 4, p, USD, inp=True)
    put(ic, r, 5, f"=D{r}/C{r}", '$0.000'); put(ic, r, 6, f"=ROUNDUP($B$4/C{r},0)", INT); put(ic, r, 7, f"=F{r}*D{r}", USD)
    put(ic, r, 8, pk, inp=True); ic.cell(r, 9, note).font = f(color=GREY)
yn(ic, "H7:H11")
ICE_COST = '=SUMIF(\'Ice & Cups\'!H7:H11,"Y",\'Ice & Cups\'!G7:G11)'

put(ic, 14, 1, "Cups needed", bold=True); put(ic, 14, 2, 250, INT, inp=True, fill=YEL)
ic.cell(14, 3, "≈ 3.5 per guest (people lose cups)").font = f(italic=True, color=GREY)
header(ic, 16, ["Clear 16 oz cups", "Store", "Per pack", "$ / pack", "$ / cup", "Packs", "Cost", "Pick?  ·  note"])
CUPS = [
 ("Member's Mark clear 16 oz", SAMS, 132, 11.28, "Y", "https://www.samsclub.com/p/members-mark-clear-plastic-cups-16oz-132ct/prod13710456"),
 ("Great Value clear 16 oz", WM, 100, 8.98, "N", "Same-day pickup. https://www.walmart.com/ip/28212763"),
 ("Choice clear PET 16 oz, case", "WebstaurantStore", 1000, 57.99, "N", "Cheapest per cup + shipping; 750 left for future MOTION nights"),
 ("H-E-B clear 18 oz", HEB, 120, 11.48, "N", "18 oz"),
 ("True clear 16 oz", "Target", 50, 16.09, "N", "Too expensive"),
]
for i, (nm, st, ct, p, pk, note) in enumerate(CUPS):
    r = 17 + i
    put(ic, r, 1, nm); put(ic, r, 2, st); put(ic, r, 3, ct, INT, inp=True); put(ic, r, 4, p, USD, inp=True)
    put(ic, r, 5, f"=D{r}/C{r}", '$0.000'); put(ic, r, 6, f"=ROUNDUP($B$14/C{r},0)", INT); put(ic, r, 7, f"=F{r}*D{r}", USD)
    put(ic, r, 8, pk, inp=True); ic.cell(r, 9, note).font = f(color=GREY)
yn(ic, "H17:H21")
CUP_COST = '=SUMIF(\'Ice & Cups\'!H17:H21,"Y",\'Ice & Cups\'!G17:G21)'

header(ic, 24, ["Other serving", "Store", "Qty", "$ each", "", "", "Cost", "Pick?  ·  note"])
OTH = [
 ("Bottled water 24-pack", HEB, 2, 3.27, "Y", "https://www.heb.com/product-detail/1042431"),
 ("Napkins 500 ct", HEB, 1, 3.48, "Y", "https://www.heb.com/product-detail/645601"),
 ("2 gal glass dispenser (one per JJ)", "Target · Threshold", 2, 29.75, "Y", "Mix each JJ in a bucket/cooler, refill the dispenser"),
 ("3 gal acrylic dispenser, ice core", "WebstaurantStore", 2, 42.99, "N", "Shatterproof, holds a whole batch"),
]
for i, (nm, st, q, p, pk, note) in enumerate(OTH):
    r = 25 + i
    put(ic, r, 1, nm); put(ic, r, 2, st); put(ic, r, 3, q, INT, inp=True); put(ic, r, 4, p, USD, inp=True)
    put(ic, r, 7, f"=C{r}*D{r}", USD); put(ic, r, 8, pk, inp=True); ic.cell(r, 9, note).font = f(color=GREY)
yn(ic, "H25:H28")
OTH_COST = '=SUMIF(\'Ice & Cups\'!H25:H28,"Y",\'Ice & Cups\'!G25:G28)'

# ---------------------------------------------------------------- Decor
dc = sheet("Decor", [34, 12, 22, 12, 34, 12, 10, 8, 50])
title(dc, "Decor · retail vs eBay / Grailed", "Best price picked per item. Grailed had nothing useful (only streetwear). Order eBay by ~12 Oct to land by 27 Oct.")
header(dc, 4, ["Item", "Zone", "Retail option", "Retail $", "eBay option", "eBay $ (+ship)", "Best $", "Pick?", "Why / link"])
DEC = [
 ("Mirror ball 20\" + motor + pin light", "Floor", "Misc Rentals, E 51st (rent)", 45.00, "EM16 16\" new", 149.99, "Y", "Rent. eBay balls 2–4× more, no used ones. miscaustin.com"),
 ("Leopard throw / rug under decks", "Floor", "Garvee 5×7 rug", 50.99, "Leopard faux-fur throw 50×60", 22.09, "Y", "Throw is $29 cheaper. ebay.com/itm/124847880165"),
 ("Gothic candelabra, 5-arm iron", "Bar", "—", None, "Vintage wrought iron 20\" (used, Best Offer)", 38.00, "Y", "No retail baseline. ebay.com/itm/158261116085"),
 ("Black LED taper candles, 4-pk", "Bar", "Walmart", 23.78, "Flameless w/ timer, ships from Houston", 14.99, "Y", "ebay.com/itm/256904542238"),
 ("Black creepy cloth (hall windows)", "Witch", "Walmart AGCAN 197×79\"", 15.99, "6-pk 30×72\" (ships from Israel, ~16 Oct)", 17.26, "Y", "Walmart piece is one big sheet — easier."),
 ("Black eucalyptus stems", "Witch", "eBay listing from 1 Oct", 30.59, "No active black listing 5 Oct", None, "Y", "Or green stems $20.69 + black spray paint"),
 ("Red filament bulbs, 2-pk", "Witch", "Home Depot · Sunlite", 14.99, "—", None, "Y", "homedepot.com"),
 ("Ornate antique-black frame 18×24", "Walls", "Target · Craig Frames", 60.99, "None that size", None, "N", "Skip unless Open Canopy art needs it"),
 ("24\" round baroque mirror, black", "Walls", "Walmart · YIDIE", 55.99, "28\" black $111.14", 111.14, "N", "Thrift first (Uncommon Objects, Top Drawer)"),
]
for i, (it, z, ro, rp, eo, ep, pk, why) in enumerate(DEC):
    r = 5 + i
    put(dc, r, 1, it); put(dc, r, 2, z); put(dc, r, 3, ro); put(dc, r, 4, rp, USD, inp=True)
    put(dc, r, 5, eo); put(dc, r, 6, ep, USD, inp=True)
    put(dc, r, 7, f'=IF(AND(D{r}="",F{r}=""),0,MIN(IF(D{r}="",1E9,D{r}),IF(F{r}="",1E9,F{r})))', USD)
    put(dc, r, 8, pk, inp=True); put(dc, r, 9, why)
dl = 5 + len(DEC) - 1
yn(dc, f"H5:H{dl}")
put(dc, dl+1, 1, "Decor (picked)", bold=True, fill=SUB); put(dc, dl+1, 7, f'=SUMIF(H5:H{dl},"Y",G5:G{dl})', USD, bold=True, fill=SUB)
DEC_COST = f"Decor!G{dl+1}"
dc.cell(dl+3, 1, "Free: borrow 2–3 pieces from local artists (Open Canopy contacts) for wall credit + IG tag + invite.").font = f(italic=True, color=GREY)
dc.cell(dl+4, 1, "eBay delivery dates were quoted to an Ohio ZIP; all active listings showed 7–17 Oct, so Austin should be similar.").font = f(italic=True, color=GREY)

# ---------------------------------------------------------------- DIY / maker
mk = sheet("DIY (Maker)", [34, 26, 26, 12, 12, 46])
title(mk, "Make it ourselves", "UT Libraries Foundry (Fine Arts Library 3.200) = free for UT students/staff, cert needed per machine. Call (512) 495-4481 for print costs.")
header(mk, 4, ["Make", "Where", "Machine", "Est. materials $", "Pick?", "Notes"])
MK = [
 ("Eye triptych, 3 × 24×36 posters", "UT Foundry", "Large-format printer", 30, "Y", "EST — Foundry print cost not posted. FedEx Guadalupe is the paid backup."),
 ("MOTION acrylic logo sign", "UT Foundry or Texas Inventionworks", "Laser cutter", 20, "Y", "EST — one 12×24 black or clear acrylic sheet."),
 ("\"VODKA\" / \"TEQUILA\" dispenser labels", "Texas Inventionworks", "Vinyl / sticker machine", 5, "Y", "EST. Bodoni, ice blue + red."),
 ("Smoking-area sign 5×7", "UT Foundry / any printer", "Printer or laser", 3, "Y", "EST. Flyer font."),
 ("Cap 01 side eyes as patches", "UT Foundry", "Baby Lock embroidery (flat)", 15, "N", "EST. No cap frame — stitch eyes flat on twill, cut, sew/heat-press on. Only if we go DIY caps."),
]
for i, (it, wh, ma, c, pk, nt) in enumerate(MK):
    r = 5 + i
    put(mk, r, 1, it); put(mk, r, 2, wh); put(mk, r, 3, ma); put(mk, r, 4, c, USD, inp=True); put(mk, r, 5, pk, inp=True); put(mk, r, 6, nt)
ml = 5 + len(MK) - 1
yn(mk, f"E5:E{ml}")
put(mk, ml+1, 1, "DIY (picked)", bold=True, fill=SUB); put(mk, ml+1, 4, f'=SUMIF(E5:E{ml},"Y",D5:D{ml})', USD, bold=True, fill=SUB)
DIY_COST = f"'DIY (Maker)'!D{ml+1}"
mk.cell(ml+3, 1, "Other spaces: Asmbly (9701 Dessau Rd, near Pflugerville) $95/mo, laser + vinyl + Brother SE600 (4×4\" max). APL Central: 3D print only.").font = f(italic=True, color=GREY)

# ---------------------------------------------------------------- Smoking
sm = sheet("Smoking", [34, 18, 8, 10, 12, 50])
title(sm, "Balcony smoking spot", "21+ only (TX). We can hand out free to 21+ guests — we can't sell. Check the Rio lease allows balcony smoking.")
header(sm, 4, ["Item", "Store", "Qty", "$ each", "Cost", "Notes"])
SM = [
 ("Cigarettes, pack (e.g. Marlboro)", "any gas station / H-E-B", 2, 9.50, "EST — Austin ≈ $9–10/pack (WPR $10.04 TX avg, Jan 2026)"),
 ("BIC Maxi lighters, 5-pk", WM, 1, 5.44, "https://www.walmart.com/ip/24705330"),
 ("Windproof lidded ashtray, black", "Walmart · XTORKIT", 1, 11.82, "Beat eBay ($31.40)"),
 ("Sand bucket for butts (backup)", "any", 1, 0, "Use a spare can + sand/water"),
]
for i, (it, st, q, p, nt) in enumerate(SM):
    r = 5 + i
    put(sm, r, 1, it); put(sm, r, 2, st); put(sm, r, 3, q, INT, inp=True); put(sm, r, 4, p, USD, inp=True)
    put(sm, r, 5, f"=C{r}*D{r}", USD); put(sm, r, 6, nt)
sl = 5 + len(SM) - 1
put(sm, sl+1, 1, "Smoking total", bold=True, fill=SUB); put(sm, sl+1, 5, f"=SUM(E5:E{sl})", USD, bold=True, fill=SUB)
SM_COST = f"Smoking!E{sl+1}"

# ---------------------------------------------------------------- Hats
h = sheet("Hats", [36, 12, 12, 12, 12, 14, 16, 50])
title(h, "Cap 01 · 24 caps, front logo + side eyes", "Blank: District DT600 distressed dad hat $6.61 (Blankstyle, $100+ cart). Need caps by Wed 28 Oct — order by ~12 Oct.")
put(h, 4, 1, "Caps", bold=True); put(h, 4, 2, 24, INT, inp=True, fill=YEL)
put(h, 5, 1, "DT600 blank $", bold=True); put(h, 5, 2, 6.61, USD, inp=True)
header(h, 7, ["Route", "Embroidery / cap", "Side / cap", "Blank / cap", "One-time $", "Total", "Per cap", "Turnaround · notes"])
HT = [
 ("Printful (Yupoong 6245CM)", 14.94, 2.95, 0, 0, "2–5 biz days + ship. Blank included. Order 25 → digitizing free; else $6.50/file. Side fee per-hat unconfirmed."),
 ("Bolt Printing (VC300A)", 13.00, 5.00, 0, 0, "9–14 biz days (tight). Front incl. digitizing ≤8k stitches; side is EST — quote."),
 ("Neatko, Leander (local)", 20.00, 20.00, None, 0, "5–7 biz days. $15–25 per placement (midpoint used). No minimum. Quote: neatko.com/quote"),
 ("Austin Print Co, Cedar Park (local)", None, None, None, 0, "Quote needed: (512) 200-2556. Digitizing incl., no setup fee, 'a few days'."),
 ("DIY patches at UT Foundry", 0, 0, None, 30, "EST materials. Labour is ours. Eyes simplified to 2–3 tones."),
]
for i, (nm, e, sd, bl, one, nt) in enumerate(HT):
    r = 8 + i
    put(h, r, 1, nm); put(h, r, 2, e, USD, inp=True); put(h, r, 3, sd, USD, inp=True)
    put(h, r, 4, "=$B$5" if bl is None else bl, USD, inp=bl is not None); put(h, r, 5, one, USD, inp=True)
    put(h, r, 6, f'=IF(B{r}="","quote",(B{r}+C{r}+D{r})*$B$4+E{r})', USD); put(h, r, 7, f'=IF(B{r}="","quote",F{r}/$B$4)', USD); put(h, r, 8, nt)
put(h, 14, 1, "Pick route (row name)", bold=True); put(h, 14, 2, "Printful (Yupoong 6245CM)", inp=True, fill=YEL)
dvh = DataValidation(type="list", formula1="=$A$8:$A$12"); h.add_data_validation(dvh); dvh.add("B14")
put(h, 15, 1, "Hats cost", bold=True); put(h, 15, 2, '=IFERROR(INDEX(F8:F12,MATCH(B14,A8:A12,0))+0,0)', USD, bold=True)
h.cell(17, 1, "Halftone eyes won't stitch clean at 2.75×1.5\" — simplify to 2–3 flat tones or bold line art before digitizing.").font = f(italic=True, color=GREY)
h.cell(18, 1, "Colourway still open: dusty rose vs plum / moss / oxblood vs black. If you sell caps, the price covers this line.").font = f(italic=True, color=GREY)
HAT_COST = "Hats!B15"

# ---------------------------------------------------------------- Budget (first sheet)
bg = wb.active; bg.title = "Budget"
for i, w in enumerate([36, 14, 54], 1): bg.column_dimensions[L(i)].width = w
bg.sheet_view.showGridLines = False
title(bg, "Devil's Night · AFTERIMAGE 001 — budget", "Fri 30 Oct 2026 · Rio · doors 9 · close 2. Prices before tax, seen 1–5 Oct 2026. Blue = inputs, green = from other tabs.")
put(bg, 4, 1, "Crew", bold=True); put(bg, 4, 2, 6, INT, inp=True)
put(bg, 5, 1, "Each puts in", bold=True); put(bg, 5, 2, 176, USD, inp=True)
put(bg, 6, 1, "Pool", bold=True); put(bg, 6, 2, "=B4*B5", USD, bold=True)
header(bg, 8, ["Line", "$", "Note"])
BL = [
 ("he/they (co.b_nguyen b2b Nioon)", 300, False, "Confirmed"),
 ("Rohan Rali", 275, False, "Confirmed 5 Oct"),
 ("Glamz", 0, False, "NOT confirmed — not budgeted"),
 ("Drinks", "=" + DR_TOTAL, True, "2 JJs + cans + soju"),
 ("Ice", ICE_COST, True, "Picked source"),
 ("Cups", CUP_COST, True, "Picked source"),
 ("Water, napkins, dispensers", OTH_COST, True, ""),
 ("Decor", "=" + DEC_COST, True, "Retail vs eBay best picks"),
 ("DIY at the maker space", "=" + DIY_COST, True, "EST materials"),
 ("Smoking spot", "=" + SM_COST, True, ""),
 ("Food", 0, False, "TBD"),
 ("CDJ backup (Rock n Roll Rentals)", 0, False, "$375/weekend only if Rapid Ric falls through"),
]
for i, (ln, v, lk, nt) in enumerate(BL):
    r = 9 + i
    put(bg, r, 1, ln); put(bg, r, 2, v, USD, inp=not lk, link=lk); put(bg, r, 3, nt)
bl = 9 + len(BL) - 1
put(bg, bl+1, 1, "Party total", bold=True, fill=SUB); put(bg, bl+1, 2, f"=SUM(B9:B{bl})", USD, bold=True, fill=SUB); bg.cell(bl+1, 3).fill = SUB
put(bg, bl+2, 1, "Left in pool", bold=True); put(bg, bl+2, 2, f"=B6-B{bl+1}", USD, bold=True)
put(bg, bl+4, 1, "Cap 01 (separate — merch)"); put(bg, bl+4, 2, "=" + HAT_COST, USD, link=True); put(bg, bl+4, 3, "Route picked on Hats tab")
put(bg, bl+5, 1, "Left if the pool also pays caps", bold=True); put(bg, bl+5, 2, f"=B{bl+2}-B{bl+4}", USD, bold=True)
put(bg, bl+7, 1, "Guest drinks covered"); put(bg, bl+7, 2, f"=Drinks!B{DR_NEED+5}", PCT, link=True); put(bg, bl+7, 3, "Our drinks vs guests × drinks each (Drinks tab)")

# ---------------------------------------------------------------- Shopping run
sr = sheet("Shopping Run", [16, 30, 70])
title(sr, "Who buys what, where, when", "Liquor stores close 9 PM and all day Sunday.")
header(sr, 4, ["When", "Where", "What"])
RUN = [
 ("This week", "eBay", "Candelabra (Best Offer $30), LED tapers $14.99, leopard throw $22.09"),
 ("This week", "Hats", "Quotes from Austin Print Co + Neatko today; Printful order of 25 as backup by ~12 Oct"),
 ("By 14 Oct", "Online (Walmart / Target)", "Creepy cloth, dispensers ×2, ashtray; Home Depot red bulbs"),
 ("By 23 Oct", "UT Foundry / Inventionworks", "Print triptych posters, cut acrylic sign, vinyl dispenser labels, smoking sign"),
 ("Week of 26 Oct", "Misc Rentals, 1901 E 51st", "Reserve + pick up 20\" mirror ball kit ($45)"),
 ("Thu 29 Oct", "Twin Liquors / WB Liquors (Costco Kelly Ln)", "Svedka 1.75, Lunazul 1.75, DeKuyper triple sec, soju"),
 ("Thu 29 Oct", "H-E-B", "Hawaiian Punch, pineapple, Sprite ×2, limeade ×2, cranberry, BeatBox, BuzzBallz ×12, White Claw 24, water, napkins, lighters"),
 ("Thu 29 Oct", "Walmart + Sam's Club", "Miller Lite 30, cups 2 × 132"),
 ("Fri 30 Oct, afternoon", "Texas Cold Front ice house, Pflugerville (24/7)", "~160 lb ice (Walmart backup). Cigarettes on the way."),
 ("Fri 30 Oct, 6 PM", "Rio", "Mix both JJs in clean buckets/coolers, chill; label dispensers"),
]
for i, row in enumerate(RUN):
    for c, v in enumerate(row, 1): put(sr, 5 + i, c, v, wrap=True)

# ---------------------------------------------------------------- Sources
so = sheet("Sources", [110])
title(so, "Sources", "Prices are listings seen on these pages 1–5 Oct 2026; re-check at checkout.")
SRC = [
 "Ice: tticlub.com/TexasColdFront · articrefresh.com/ice-vending-machine · twinliquors.com (20 lb ice) · reddyice.com/locations/austin-301",
 "Cups: samsclub.com prod13710456 · walmart.com/ip/28212763 · webstaurantstore.com 500CC16 · heb.com 2812012 · target.com A-90367920",
 "eBay: itm 158261116085, 124847880165, 256904542238, 297101089464, 306980432913, 166387828199, 188940592615, 198665930673",
 "Hats: blankstyle.com (DT600, 6245CM) · help.printful.com (digitization fees, fulfillment time) · stylefactoryproductions.com/blog/printful-pricing · boltprinting.com · neatko.com · austinprintco.com/custom-hats-austin · logostore.org",
 "Maker: lib.utexas.edu/visit-us/library-spaces/foundry · guides.lib.utexas.edu/Foundry/SewingAndTextiles · inventionworks.engr.utexas.edu · asmbly.org/space",
 "Smoking: comptroller.texas.gov/taxes/tobacco/faq-reg.php · worldpopulationreview.com/state-rankings/cigarette-prices-by-state · walmart.com/ip/24705330",
 "Drinks: twinliquors.com (Svedka, Lunazul, DeKuyper, soju) · heb.com (BeatBox 5508020, BuzzBallz 2787503, White Claw 3758568, cranberry 5785031) · walmart.com/ip/10984520 (Miller Lite)",
]
for i, s_ in enumerate(SRC): put(so, 4 + i, 1, s_, wrap=True)

for s in wb.worksheets:
    s.freeze_panes = None
wb.save("Devils_Night_Budget_and_Shopping.xlsx")
print("ok")
